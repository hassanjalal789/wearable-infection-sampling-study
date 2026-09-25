from __future__ import annotations

import json
import math
import statistics
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results"
RESULTS.mkdir(parents=True, exist_ok=True)

sys.path.insert(0, str(ROOT / "src"))
from gate_reproduction import gate_b_report  # noqa: E402


def _norm_text(v) -> str:
    if v is None:
        return ""
    return " ".join(str(v).strip().split())


def _num(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        if isinstance(v, float) and math.isnan(v):
            return None
        return float(v)
    s = _norm_text(v)
    if s.lower() in {"", "nan", "null", "na", "miss", "missed"}:
        return None
    try:
        return float(s)
    except ValueError:
        return None


def _mishra() -> dict:
    path = ROOT / "metadata" / "phase1_supplementary_data.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb["SuppTable13_Fig4abc"]

    covid_rows = []
    in_covid = False

    for row in ws.iter_rows(values_only=True):
        first = _norm_text(row[0] if len(row) > 0 else None)
        second = _norm_text(row[1] if len(row) > 1 else None)

        if first == "ParticipantID" and second == "Category":
            in_covid = True
            continue

        if in_covid and first.lower().startswith("for figure 4c"):
            break

        if not in_covid or not first:
            continue

        category = second.lower()
        if category not in {"single region", "early & multiple region", "others"}:
            continue

        covid_rows.append({
            "participant_id": first,
            "category": category,
            "symptom_delta_raw": row[2] if len(row) > 2 else None,
            "diagnosis_delta_raw": row[3] if len(row) > 3 else None,
        })

    detected = [
        r for r in covid_rows
        if r["category"] in {"single region", "early & multiple region"}
    ]

    symptom_info = []
    symptom_numeric = []

    for r in detected:
        raw = _norm_text(r["symptom_delta_raw"])
        unavailable = raw.lower() in {"", "nan", "no symptom day"}
        if not unavailable:
            symptom_info.append(r)

        x = _num(r["symptom_delta_raw"])
        if x is not None:
            symptom_numeric.append(x)

    at_or_before = sum(x <= 0 for x in symptom_numeric)
    median_delta = statistics.median(symptom_numeric)
    median_lead = -float(median_delta)

    return {
        "analysed_cases": len(covid_rows),
        "detected_cases_groups_I_II": len(detected),
        "symptom_information_cases_among_detected": len(symptom_info),
        "rhrdiff_numeric_symptom_deltas": len(symptom_numeric),
        "detected_at_or_before_onset": at_or_before,
        "median_lead_days_vs_onset": median_lead,
        "median_signed_delta_days": float(median_delta),
        "participant_ids": [r["participant_id"] for r in covid_rows],
        "source": "phase1_supplementary_data.xlsx::SuppTable13_Fig4abc",
        "interpretation": (
            "Groups I and II are the 26 detected cases described in the paper. "
            "'NaN' denotes unavailable symptom/diagnosis date; 'miss' denotes an "
            "algorithm miss and is therefore not converted into a numeric RHR-Diff delta."
        ),
    }


def _alavi() -> dict:
    path = ROOT / "metadata" / "phase2_covid_metadata.xlsx"
    wb = load_workbook(path, read_only=True, data_only=True)

    pos = wb["SourceData_COVID19_Positives"]
    fig4a = wb["SourceData_Fig4A_Presymp"]
    fig4b = wb["SourceData_Fig4B_Asymp"]

    positives = []
    for i, row in enumerate(pos.iter_rows(values_only=True), start=1):
        if i == 1:
            continue
        pid = _norm_text(row[0] if row else None)
        if pid:
            positives.append(pid)

    def extract_alert_sheet(ws):
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i == 1:
                continue
            pid = _norm_text(row[0] if len(row) > 0 else None)
            if not pid:
                continue
            raw = row[1] if len(row) > 1 else None
            rows.append((pid, _num(raw), _norm_text(raw)))
        return rows

    symptomatic = extract_alert_sheet(fig4a)
    asymptomatic = extract_alert_sheet(fig4b)

    alert_ids = [p for p, _, _ in symptomatic + asymptomatic]
    duplicate_alert_ids = sorted({p for p in alert_ids if alert_ids.count(p) > 1})
    same_positive_set = set(alert_ids) == set(positives)

    sym_tp = sum(v is not None and v <= 0 for _, v, _ in symptomatic)
    asym_tp = sum(v is not None and v <= 0 for _, v, _ in asymptomatic)
    tp = sym_tp + asym_tp
    sensitivity = 100.0 * tp / len(positives)

    return {
        "confirmed_positives": len(positives),
        "symptomatic_rows": len(symptomatic),
        "asymptomatic_rows": len(asymptomatic),
        "symptomatic_at_or_before_onset": sym_tp,
        "asymptomatic_at_or_before_test": asym_tp,
        "true_positive_participants": tp,
        "sensitivity_per_person_pct": sensitivity,
        "fig4_union_matches_positive_cohort": same_positive_set,
        "duplicate_alert_ids": duplicate_alert_ids,
        "source": (
            "phase2_covid_metadata.xlsx::"
            "SourceData_COVID19_Positives + SourceData_Fig4A_Presymp + "
            "SourceData_Fig4B_Asymp"
        ),
    }


def compute_gate_b() -> dict:
    mishra = _mishra()
    alavi = _alavi()

    observed = {
        "mishra_analysed_cases": mishra["analysed_cases"],
        "mishra_detected": mishra["detected_cases_groups_I_II"],
        "mishra_median_lead_days_vs_onset": mishra["median_lead_days_vs_onset"],
        "mishra_detected_at_or_before_onset_of_25": mishra["detected_at_or_before_onset"],
        "alavi_confirmed_positives": alavi["confirmed_positives"],
        "alavi_sensitivity_per_person_pct": alavi["sensitivity_per_person_pct"],
    }

    report = gate_b_report(observed)

    report["alavi_specificity_per_alert_day_pct"] = {
        "status": "NOT_RECONSTRUCTIBLE_PUBLIC_DATA",
        "published": 87.7,
        "tol": 2.0,
        "published_numerator_TN": 87124,
        "published_denominator_TN_plus_FP": 87124 + 12186,
        "published_false_positive_alert_days": 12186,
        "published_arithmetic_pct": 100.0 * 87124 / (87124 + 12186),
        "reason": (
            "The public Source Data workbook supplied here covers COVID-positive "
            "participants and alert timing, but does not provide the non-COVID "
            "participant-day alert labels and negative-test/untested survey timing "
            "needed to independently reconstruct TN=87,124 and FP=12,186. "
            "The values are retained as a paper-reported benchmark, not counted as "
            "a reproduced statistic."
        ),
    }

    return {
        "gate": "B",
        "status": "COMPLETE_WITH_ONE_PUBLIC_DATA_LIMITATION",
        "observed": observed,
        "mishra_detail": mishra,
        "alavi_detail": alavi,
        "report": report,
        "decision": (
            "Six preregistered quantities are independently reconstructed from "
            "participant-level public source tables. Alavi alert-day specificity "
            "is not labelled PASS or FAIL because the public source package used "
            "here lacks the denominator-level non-COVID alert/survey labels needed "
            "for independent reconstruction."
        ),
    }


def main():
    result = compute_gate_b()
    out = RESULTS / "gate_b_results.json"
    out.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
